"""Genera el Excel de evaluación humana de POLARIDAD (gold v2 · 2026-06-12).

Lecciones aplicadas del ejercicio de abril (evaluations/2026-04-18/human/REPORTE-ACUERDO-HUMANO.md):
- Clave estable: platform_comment_id (sobrevive re-ingests), NUNCA PK de BD.
- Campo evaluado: POLARIDAD 3 clases (no tono-11, donde el techo humano es ~50%).
- Opciones de ELECCIÓN (dropdown), no texto libre + escape "No claro" para no forzar adivinanza.
- Contexto visible: dirigente, rol, post al que responde (la polaridad depende de la entidad).
- Hoja Calibración con 5 ejemplos resueltos (mitiga divergencia tipo h3).
- Label del LLM OCULTO en las copias de evaluadores; vive solo en el master del repo.

Salida: backend/evaluations/2026-06-12-gold-polaridad/
  master_oculto.csv (CON label LLM · NO se envía) + evaluador_{1,2,3}.xlsx (idénticos)

Uso (HOST, crece-db arriba):
  PYTHONPATH=. .venv/bin/python scripts/gen_gold_excel_polaridad.py
"""

from __future__ import annotations

import random
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 42
N_TOTAL = 100  # ~33 por clase de polaridad LLM (oculta)
OUT = Path(__file__).resolve().parent.parent / "evaluations" / "2026-06-12-gold-polaridad"

OPCIONES = '"Positivo para el dirigente,Neutro,Negativo para el dirigente,No claro / no se entiende"'

CALIBRACION = [
    ("¡Felicidades! Oaxaca está en buenas manos con usted 💪", "oficialismo",
     "Positivo para el dirigente", "Elogio directo al dirigente."),
    ("Puro show, mejor pónganse a trabajar", "oficialismo",
     "Negativo para el dirigente", "Crítica dirigida al dirigente/su gestión."),
    ("¿Alguien sabe a qué hora empieza el evento?", "oficialismo",
     "Neutro", "Pregunta logística; no evalúa al dirigente."),
    ("Este gobierno son unos rateros, ya queremos cambio", "oposicion",
     "Positivo para el dirigente", "OJO: ataca al GOBIERNO; si el dirigente es OPOSICIÓN, "
     "el comentario juega A FAVOR de su causa. La columna 'rol' importa."),
    ("jajaja 😂😂", "oficialismo",
     "No claro / no se entiende", "Sin contexto suficiente para decidir — usar el escape, no adivinar."),
]


def sql(q: str) -> list[str]:
    r = subprocess.run(
        ["docker", "exec", "crece-db", "psql", "-U", "crece", "-d", "crece", "-tA", "-F", "\t", "-c", q],
        capture_output=True, text=True,
    )
    if r.returncode != 0:
        sys.exit(f"SQL error: {r.stderr[:300]}")
    return [l for l in r.stdout.splitlines() if l.strip()]


def fetch_sample() -> list[dict]:
    rows: list[dict] = []
    # ~equilibrado por clase LLM (oculta) · dentro de clase, mezcla de dirigentes (seed fijo)
    for pol, n in (("1", 33), ("0", 33), ("-1", 34)):
        q = f"""
        SELECT c.platform_comment_id, d.full_name, d.rol_politico, sp.platform,
               left(replace(replace(c.content, E'\\t',' '), E'\\n',' '), 400),
               left(replace(replace(p.content, E'\\t',' '), E'\\n',' '), 200),
               c.nlp_polaridad
        FROM social_comments c
        JOIN social_posts p ON p.id = c.parent_post_id
        JOIN social_profiles sp ON sp.id = p.profile_id
        JOIN dirigentes d ON d.id = sp.dirigente_id
        WHERE c.nlp_polaridad = {pol} AND length(trim(c.content)) >= 8
          AND d.id <> 56  -- fuera RSS bot
        ORDER BY md5(c.platform_comment_id || '{SEED}')
        LIMIT {n}
        """
        for line in sql(q):
            f = line.split("\t")
            if len(f) < 7:
                continue
            rows.append({
                "key": f[0], "dirigente": f[1], "rol": f[2] or "?", "plataforma": f[3],
                "comentario": f[4], "post_contexto": f[5], "llm_polaridad": f[6],
            })
    random.Random(SEED).shuffle(rows)  # que el orden no delate la clase
    return rows


def build_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    # ── Hoja 1 · Instrucciones (simples) ──
    ws = wb.active
    ws.title = "Instrucciones"
    ws.column_dimensions["A"].width = 100
    lineas = [
        ("Evaluación de comentarios — ¿el comentario es positivo, neutro o negativo PARA el dirigente?", True),
        ("", False),
        ("1. Ve a la hoja 'Clasificar'. Hay 100 comentarios.", False),
        ("2. En la columna TU RESPUESTA elige UNA opción del menú (flecha al seleccionar la celda).", False),
        ("3. Las opciones son 4: Positivo para el dirigente · Neutro · Negativo para el dirigente · No claro.", False),
        ("4. Usa las columnas de contexto: quién es el dirigente, su rol (oficialismo/oposición) y el post original.", False),
        ("   CLAVE: un comentario que ataca al gobierno es POSITIVO para un dirigente de OPOSICIÓN.", False),
        ("5. Si de verdad no se entiende, elige 'No claro / no se entiende' — NO adivines.", False),
        ("6. Antes de empezar, revisa la hoja 'Calibración' (5 ejemplos ya resueltos, 2 minutos).", False),
        ("7. No hay respuestas correctas 'oficiales' — queremos TU criterio. Tiempo estimado: 30-40 min.", False),
    ]
    for i, (txt, b) in enumerate(lineas, 1):
        ws.cell(row=i, column=1, value=txt).font = Font(bold=b)

    # ── Hoja 2 · Calibración ──
    ws = wb.create_sheet("Calibración")
    for col, (w, h) in zip("ABCD", [(50, "comentario"), (14, "rol del dirigente"), (28, "respuesta correcta"), (52, "por qué")]):
        ws.column_dimensions[col].width = w
        c = ws.cell(row=1, column=ord(col) - 64, value=h)
        c.font = bold; c.fill = head_fill
    for i, (txt, rol, resp, why) in enumerate(CALIBRACION, 2):
        for j, v in enumerate((txt, rol, resp, why), 1):
            ws.cell(row=i, column=j, value=v).alignment = wrap

    # ── Hoja 3 · Clasificar ──
    ws = wb.create_sheet("Clasificar")
    headers = ["#", "clave (no tocar)", "dirigente", "rol", "plataforma", "post original (contexto)",
               "COMENTARIO A EVALUAR", "TU RESPUESTA", "notas (opcional)"]
    widths = [5, 24, 22, 13, 11, 38, 55, 30, 25]
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = bold; c.fill = head_fill
        ws.column_dimensions[chr(64 + j)].width = w
    dv = DataValidation(type="list", formula1=OPCIONES, allow_blank=True, showDropDown=False)
    dv.error = "Elige una opción del menú"; dv.errorStyle = "stop"
    ws.add_data_validation(dv)
    for i, r in enumerate(rows, 2):
        vals = [i - 1, r["key"], r["dirigente"], r["rol"], r["plataforma"],
                r["post_contexto"], r["comentario"], "", ""]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v).alignment = wrap
        dv.add(ws.cell(row=i, column=8))
    ws.freeze_panes = "A2"
    wb.save(path)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    rows = fetch_sample()
    print(f"muestra: {len(rows)} comments · dirigentes: {len({r['dirigente'] for r in rows})}")
    # master con label oculto (queda en repo, NO se envía)
    import csv
    with open(OUT / "master_oculto.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)
    # 3 copias idénticas sin label
    for i in (1, 2, 3):
        build_xlsx(rows, OUT / f"evaluador_{i}.xlsx")
    print(f"OK → {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
